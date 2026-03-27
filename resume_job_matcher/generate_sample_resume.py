"""Generate a sample resume as a DOCX file for testing the Resume Job Matcher."""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os


def generate_sample_resume(output_path="uploads/sample_resume.docx"):
    doc = Document()

    # Name
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run("John Smith")
    run.bold = True
    run.font.size = Pt(20)

    # Contact info
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.add_run("john.smith@email.com | (555) 123-4567 | San Francisco, CA | linkedin.com/in/johnsmith")

    # Summary
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(
        "Senior Full Stack Developer with 10+ years of experience specializing in backend development. "
        "Expertise in C#, .NET, ASP.NET Core, Python, JavaScript, and cloud technologies. "
        "Proven track record of designing and building scalable enterprise applications, REST APIs, "
        "microservices architectures, and distributed systems. Strong focus on backend performance, "
        "database optimization, and system reliability. Experienced in leading cross-functional teams "
        "and mentoring junior developers in agile environments."
    )

    # Skills
    doc.add_heading("Technical Skills", level=1)
    skills = doc.add_paragraph()
    skills.add_run("Languages: ").bold = True
    skills.add_run("C#, Python, JavaScript, TypeScript, SQL, HTML, CSS, Java, Go\n")
    skills.add_run("Backend Frameworks: ").bold = True
    skills.add_run("ASP.NET Core, .NET Framework, Entity Framework, Django, Flask, FastAPI, Spring\n")
    skills.add_run("Frontend: ").bold = True
    skills.add_run("React, Angular, Vue, Bootstrap, Tailwind\n")
    skills.add_run("Databases: ").bold = True
    skills.add_run("SQL Server, PostgreSQL, MongoDB, Redis, Elasticsearch, DynamoDB\n")
    skills.add_run("Cloud & DevOps: ").bold = True
    skills.add_run("Azure, AWS, Docker, Kubernetes, Terraform, Jenkins, GitHub Actions, CI/CD\n")
    skills.add_run("Tools & Practices: ").bold = True
    skills.add_run("Git, Jira, Agile, Scrum, Microservices, REST API, GraphQL, System Design, "
                   "Unit Testing, Integration Testing, Test Driven Development\n")
    skills.add_run("Soft Skills: ").bold = True
    skills.add_run("Leadership, Project Management, Communication, Mentoring, Problem Solving, "
                   "Stakeholder Management")

    # Experience
    doc.add_heading("Professional Experience", level=1)

    # Job 1
    job1_title = doc.add_paragraph()
    run = job1_title.add_run("Senior Software Engineer at TechGlobal Solutions")
    run.bold = True
    doc.add_paragraph("January 2020 - Present")
    bullets1 = [
        "Lead a team of 8 developers building enterprise-grade .NET microservices handling 50M+ daily transactions",
        "Architected migration from monolithic ASP.NET MVC to ASP.NET Core microservices, reducing deployment time by 70%",
        "Designed and implemented RESTful APIs consumed by 200+ internal and external clients",
        "Optimized SQL Server database queries resulting in 40% performance improvement",
        "Implemented CI/CD pipelines using Azure DevOps and Docker/Kubernetes for automated deployments",
        "Mentored 5 junior developers and conducted code reviews ensuring code quality standards",
    ]
    for b in bullets1:
        doc.add_paragraph(b, style="List Bullet")

    # Job 2
    job2_title = doc.add_paragraph()
    run = job2_title.add_run("Full Stack Developer at InnovateTech Inc.")
    run.bold = True
    doc.add_paragraph("March 2017 - December 2019")
    bullets2 = [
        "Built and maintained full-stack applications using C#, ASP.NET Core, React, and SQL Server",
        "Developed RESTful APIs and background services processing financial data for banking clients",
        "Implemented Redis caching layer reducing API response times by 60%",
        "Collaborated with product managers to translate business requirements into technical solutions",
        "Led migration of legacy systems to cloud-based Azure infrastructure",
    ]
    for b in bullets2:
        doc.add_paragraph(b, style="List Bullet")

    # Job 3
    job3_title = doc.add_paragraph()
    run = job3_title.add_run("Software Developer at CodeWorks LLC")
    run.bold = True
    doc.add_paragraph("June 2014 - February 2017")
    bullets3 = [
        "Developed web applications using C#, .NET Framework, JavaScript, and SQL Server",
        "Built automated testing suites with NUnit and integration tests reducing bug rates by 30%",
        "Worked on payment processing systems handling PCI-compliant financial transactions",
        "Participated in agile sprints, daily standups, and sprint retrospectives",
    ]
    for b in bullets3:
        doc.add_paragraph(b, style="List Bullet")

    # Job 4
    job4_title = doc.add_paragraph()
    run = job4_title.add_run("Junior Developer at StartupForge")
    run.bold = True
    doc.add_paragraph("August 2012 - May 2014")
    bullets4 = [
        "Developed features for SaaS platform using C#, ASP.NET MVC, and jQuery",
        "Managed PostgreSQL databases and wrote stored procedures for reporting",
        "Implemented responsive frontend components using HTML, CSS, and JavaScript",
    ]
    for b in bullets4:
        doc.add_paragraph(b, style="List Bullet")

    # Education
    doc.add_heading("Education", level=1)
    edu = doc.add_paragraph()
    run = edu.add_run("Bachelor's in Computer Science")
    run.bold = True
    doc.add_paragraph("University of California, Berkeley - 2012")

    # Certifications
    doc.add_heading("Certifications", level=1)
    certs = [
        "Microsoft Certified: Azure Developer Associate",
        "AWS Certified Solutions Architect - Associate",
        "Certified Scrum Master (CSM)",
    ]
    for c in certs:
        doc.add_paragraph(c, style="List Bullet")

    # Save
    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    doc.save(output_path)
    print(f"Sample resume saved to: {output_path}")
    return output_path


if __name__ == "__main__":
    generate_sample_resume()
